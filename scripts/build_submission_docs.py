"""Build the revised report, artifact guide, and code PDF from verified project facts."""

from pathlib import Path
import csv
import textwrap

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission"
OUT.mkdir(exist_ok=True)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.78)
    sec.left_margin = sec.right_margin = Inches(0.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in [("Title", 18, 0, 10), ("Heading 1", 13, 12, 5), ("Heading 2", 11, 9, 4)]:
        st = styles[name]
        st.font.name = "Aptos Display" if name == "Title" else "Aptos"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ["Subtitle"]:
        styles[name].font.name = "Aptos"
        styles[name].font.size = Pt(10)
        styles[name].font.color.rgb = RGBColor(70, 70, 70)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Group 5 | Metro Manila AADT simulation")
    footer.runs[0].font.size = Pt(8)


def title(doc, text, subtitle):
    doc.add_paragraph(text, "Title")
    doc.add_paragraph(subtitle, "Subtitle")


def para(doc, text):
    doc.add_paragraph(text)


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    tcpr.append(borders)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for row in rows:
        cells = t.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_border(cell)
            shade(cell, "DCE8F1" if i == 0 else ("F6F8FA" if i % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
                    if i == 0:
                        run.font.bold = True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph("")


def project_paper():
    doc = Document()
    style_document(doc)
    title(doc, "Metro Manila AADT Traffic Risk Simulation", "Revised project paper aligned with the working Python artifact | Group 5")
    doc.add_heading("Abstract", 1)
    para(doc, "This project models annual traffic demand for 20 Metro Manila roads from a 2025 baseline through 2035. Vehicle-class counts from a 2020-2025 CSV are converted to passenger-car units (PCU), advanced under stated scenarios, and compared with scenario road capacities. Risk is reported as Low, Medium, or High using volume-to-capacity thresholds. A Random Forest independently labels the same risk as an AI comparison. The road capacities are synthetic, source AADT rows remain unverified, and the results should be read as sensitivity analysis rather than measured forecasts or physical failure dates.")

    doc.add_heading("1 Introduction and objective", 1)
    para(doc, "The objective is to show which included roads may reach high use of assumed capacity first under different annual traffic and intervention assumptions. The annual time step matches the dataset's AADT granularity. The model supports discussion and review; it does not represent hourly queues, intersection operations, or automatic policy decisions.")

    doc.add_heading("2 Dataset and preparation", 1)
    para(doc, "The supplied dataset contains 120 road-year rows: 20 named roads in each year from 2020 through 2025. The road-type field contains 6 Circumferential, 12 Radial, and 2 Highway roads. The 2025 rows report 3,959,637 vehicles per day in total across the 20 road records. This is a sum of corridor counts, not a count of unique trips across Metro Manila. Forty historical rows are flagged as COVID-19 pandemic years. The CSV marks 114 AADT rows as unverified against source PDFs and 6 as requiring source checks; its PCU equivalency profile is also marked unverified.")
    para(doc, "The code validates required fields, one row per road and year, nonnegative class counts, and positive capacity. It recomputes PCU as the sum of each vehicle-class count multiplied by that class's supplied PCEF. In the CSV the factors are: car 1.0, PUJ 1.4, UV 1.2, taxi 1.0, PUB 2.5, truck 2.5, trailer 2.2, motorcycle 0.5, and tricycle 1.2. These are project input assumptions pending source validation. The recomputed PCU agrees with the stored PCU within floating-point rounding.")
    para(doc, "The capacity field is explicitly labeled SYNTHETIC_SCENARIO_NOT_MEASURED. It was calibrated from 2025 demand and assumed use of capacity; it is not an observed engineering capacity. Consequently, every V/C risk result depends on this assumed denominator. The stored 2025 risk label for R-10 Del Pan is Medium because a displayed V/C rounds to 0.70. The full-precision ratio is 0.699999936, so the simulation's threshold rule labels it Low. The implementation uses full precision consistently.")

    doc.add_heading("3 Simulation design", 1)
    para(doc, "The model is a macroscopic annual time-step simulation implemented with pandas and NumPy. For each road it stores the current year, vehicle-class volumes, recomputed road PCU volume, scenario capacity limit, threshold risk, Random Forest risk, and the first year the threshold risk becomes High. It records the 2025 baseline unchanged, then compounds each vehicle class once per year for 2026-2035. The baseline growth assumption is 4.5% annually and baseline capacity is fixed. This uniform class growth keeps the vehicle mix constant unless a named scenario changes it.")
    para(doc, "For each road-year, V/C equals PCU volume divided by scenario capacity. Low means V/C below 0.70; Medium means 0.70 to below 0.95; High means 0.95 or above. The first High year records the first threshold crossing within 2025-2035, including a High 2025 baseline. It is not a validated road failure date. The terminal's group alert is the highest risk among the 20 roads, with counts shown for each class; it does not mean every road has that same risk.")

    doc.add_heading("4 AI integration and evaluation", 1)
    para(doc, "A scikit-learn RandomForestClassifier is trained on 2020-2023 rows (80 examples) and evaluated on 2024-2025 rows (40 examples). Features are road PCU volume, scenario capacity, year, road type, and vehicle-class shares. The forest uses 300 trees, maximum depth 8, minimum leaf size 2, balanced class weights, and random seed 42. The code displays its prediction as AI risk beside the threshold-based Risk_Level. The first High year and group alert use the threshold risk, not the forest prediction.")
    table(doc, ["Holdout measure", "Observed result"], [
        ("Accuracy", "72.5% (29 of 40)"),
        ("Macro precision", "71.9%"),
        ("Macro recall", "74.5%"),
        ("Macro F1", "72.1%"),
        ("Forest versus threshold", "11 disagreements in 40 holdout rows"),
    ], [2.2, 4.4])
    para(doc, "The paper's previous 80% accuracy figure was a target, not an achieved result. The observed holdout accuracy is below that target. Because the target labels are defined from PCU/capacity and the same quantities enter the model, the score largely tests whether the forest recovers the threshold rule; it is not independent traffic-forecast accuracy. Tree predictions may also plateau when simulated future volumes exceed the historical training range. The program therefore retains the direct rule as the interpretable risk assessment and shows AI disagreements rather than hiding them.")

    doc.add_heading("5 Scenarios and preliminary results", 1)
    table(doc, ["Scenario", "Change to the 2025 baseline"], [
        ("Baseline", "4.5% annual growth; fixed synthetic capacity"),
        ("Accelerated", "7.0% annual growth; fixed synthetic capacity"),
        ("Transit shift", "4.5% growth; from 2026, 20% of projected cars transfer to public buses, assuming one additional bus per 40 shifted cars"),
        ("EDSA and C-5 expansion", "4.5% growth; 20% more synthetic capacity on C-4 EDSA and C-5 CP Garcia from 2028"),
        ("Custom CLI", "User-selected annual growth from 0% to 30%; fixed synthetic capacity"),
    ], [1.7, 4.9])
    table(doc, ["Scenario", "2025 High", "2026 High", "2035 High"], [
        ("Baseline 4.5%", "3", "5", "20"),
        ("Accelerated 7%", "3", "6", "20"),
        ("Transit shift 20%", "3", "0", "18"),
        ("EDSA and C-5 expansion", "3", "5", "20"),
    ], [2.7, 1.3, 1.3, 1.3])
    para(doc, "These counts are from the current four-scenario output CSV and use the full-precision threshold rule. The expansion has no effect in 2026 because it begins in 2028. The transit result depends strongly on its illustrative 40-car-per-bus conversion. By 2035 all baseline roads reach High against their synthetic capacity, which highlights sensitivity to the assumed denominator and uniform compounding rather than a verified regional outcome.")

    doc.add_heading("6 Limits and responsible interpretation", 1)
    for item in [
        "The 20 selected corridors do not represent every street or unique trip in Metro Manila.",
        "AADT cannot describe peak hours, daily variation, floods, incidents, lane restrictions, or intersection delay.",
        "The 2020-2021 pandemic flags are displayed but not adjusted out of model training; a policy study should test sensitivity to excluding them.",
        "Capacity, PCEF factors, source AADT, and the 4.5% growth assumption require independent verification before operational use.",
        "Scenario variation is not a calibrated probability interval. The 2025-2035 first High year is conditional on the chosen assumptions.",
    ]:
        bullet(doc, item)
    para(doc, "The project is a decision-support exercise. Human reviewers remain responsible for source verification, interpretation, and any intervention decisions. The original draft includes an AI-use log and member-contribution claims; the group should confirm those records against its actual work before final submission.")

    doc.add_heading("7 Changes made from the earlier draft", 1)
    table(doc, ["Earlier draft statement", "Correction aligned to the artifact"], [
        ("Twenty arterial roads described as only Radial and Circumferential", "Include 2 Highway records and state the 6/12/2 type split."),
        ("Capacity and PCU factors presented as validated standards", "Label the CSV values as synthetic or unverified project inputs."),
        ("Random Forest assigns the reported Risk_Level and failure year", "Risk_Level and first High year use the direct V/C rule; RF_Risk_Level is shown separately."),
        ("At least 80% accuracy implied", "Report observed 72.5% holdout accuracy and retain 80% only as a target."),
        ("Failure year and policy trigger", "Use conditional first High threshold-crossing year; no automatic policy trigger."),
        ("Preliminary results left blank and artifact described as unfinished", "Provide executed scenario counts, source files, CLI instructions, and sample output."),
        ("2025 total stated as 3,939,637", "Correct the sum of reported 2025 road counts to 3,959,637."),
    ], [2.45, 4.15])
    doc.save(OUT / "revised_project_paper.docx")


def artifact_guide():
    doc = Document()
    style_document(doc)
    title(doc, "Working Traffic Simulation Artifact", "Group 5 submission guide and execution evidence")
    doc.add_heading("Purpose and current status", 1)
    para(doc, "The Python artifact runs an annual 2025-2035 AADT scenario for 20 Metro Manila roads and prints risk results in the terminal. It validates the dataset, converts vehicle counts to PCU, advances road state, applies a Random Forest classifier, and exports CSV results. The central result is a per-road Low, Medium, or High threshold risk for the selected year. Road capacity is synthetic, so the output is a scenario exercise.")

    doc.add_heading("Submission file inventory", 1)
    table(doc, ["File", "Purpose"], [
        ("run_simulation.py", "Interactive terminal entry point: prompts, selected-year summary, risk display, CSV export"),
        ("aadt_model.py", "Data validation, PCU conversion, Random Forest evaluation, annual scenario engine"),
        ("python_source_code.pdf", "Readable PDF copy of both Python source files"),
        ("metro_manila_aadt_simulation_input_2020_2025.csv", "Required historical dataset, 120 road-year rows"),
        ("requirements-cli.txt", "Terminal dependencies: NumPy, pandas, scikit-learn"),
        ("README.md", "Run instructions and scenario descriptions"),
        ("sample_road_year_forecast.csv", "Sample output from four predefined scenarios"),
        ("sample_2035_summary.csv", "Compact sample 2035 output"),
    ], [2.7, 3.9])

    doc.add_heading("How to run the terminal simulation", 1)
    para(doc, "Open PowerShell in the project folder. Use Python 3.12 or a compatible Python 3 installation. Run these commands in order:")
    for cmd in ["py -m venv .venv", r".\.venv\Scripts\python -m pip install -r requirements-cli.txt", r".\.venv\Scripts\python run_simulation.py"]:
        p = doc.add_paragraph(cmd)
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
    para(doc, "The program asks for scenario 1-5, then a year from 2025 through 2035; pressing Enter selects 2035. After the result box, answer y to export the entire scenario to outputs/terminal/<scenario>/ or n to return to the menu. Enter 0 at the menu to quit. The relative CSV path assumes the supplied project folder structure remains intact. If Python's Windows launcher is unavailable, install Python or use the full path to a compatible Python executable.")

    doc.add_heading("Core workflow and state updates", 1)
    for item in [
        "Load the 2020-2025 CSV and check its columns, road-year coverage, counts, and capacities.",
        "Recompute daily PCU from nine vehicle classes and the supplied PCEF values.",
        "Train the Random Forest on 2020-2023 and show its 2024-2025 holdout evaluation.",
        "Initialize each road from its 2025 counts, PCU volume, and synthetic capacity.",
        "For 2026-2035, apply the chosen annual growth and intervention assumptions, then recompute PCU and V/C.",
        "Calculate direct threshold risk, collect the separate Random Forest risk, and record the first High year.",
        "Display the selected-year risk box and optionally export the full annual result table.",
    ]:
        bullet(doc, item)
    para(doc, "The stored state includes Year, per-class vehicle counts, Road_PCU_Volume, Capacity_Limit, Risk_Level, RF_Risk_Level, and First_High_Year. The 2025 row is retained as baseline; the first growth step occurs in 2026. The terminal group's selected-year alert is the maximum per-road threshold risk, with class counts shown to avoid implying every road has the same label.")

    doc.add_heading("AI component and scenario demonstration", 1)
    para(doc, "The Random Forest has 300 trees, maximum depth 8, minimum leaf size 2, balanced class weights, and random seed 42. It consumes PCU volume, capacity, year, road type, and vehicle-class shares. On the 2024-2025 holdout it matched 72.5% of derived risk labels (29/40); macro F1 was 72.1%. Because its labels are based on the same PCU/capacity inputs, this is not independent traffic prediction. The terminal shows AI risk beside the direct threshold risk and counts disagreements.")
    para(doc, "Scenario 1 is an executable baseline with 4.5% compounded annual class growth and fixed capacity. Additional terminal choices run 7% growth, a 20% car-to-bus transfer using one added bus per 40 shifted cars, a 20% synthetic capacity increase on EDSA and C-5 from 2028, or a custom annual growth percentage. These are illustrative experiments, not observed interventions.")

    doc.add_heading("Sample output and how to read it", 1)
    para(doc, "A baseline run for 2026 displays a High group alert because 5 of 20 roads are High, 12 Medium, and 3 Low. The group alert uses the highest road risk. The road table shows the supplied 2025 raw vehicle count, calculated 2025 PCU, projected 2026 PCU, percentage of assumed capacity used, direct risk, AI risk, and first High year. The CSV exports include one row per road per year.")
    table(doc, ["Road in 2026 baseline", "Capacity used", "Threshold risk", "AI risk"], [
        ("C-4 EDSA", "104.5%", "High", "High"),
        ("R-7 Commonwealth", "102.4%", "High", "High"),
        ("R-3 SSH", "100.3%", "High", "Medium"),
    ], [2.5, 1.3, 1.4, 1.4])
    para(doc, "The R-3 SSH row illustrates why the two risk columns are separate: its V/C exceeds the High threshold while the forest returns Medium. The threshold label is the reported risk and controls the first High year. A 2025 baseline with V/C already at or above 0.95 can have First_High_Year = 2025.")

    doc.add_heading("Validation and limitations", 1)
    para(doc, "The working run produced 220 rows per predefined scenario (20 roads multiplied by 11 years), without duplicate road-year combinations. Validation checks cover the 2025 baseline, compound growth, fixed baseline capacity, and exact threshold boundaries at 0.70 and 0.95. The 2025 risk distribution recomputed from full-precision PCU/capacity is 5 Low, 12 Medium, and 3 High. The stored CSV label for R-10 Del Pan differs because its displayed 0.70 ratio is rounded from 0.699999936.")
    para(doc, "The AADT source rows and PCU factors remain unverified, capacity is synthetic, and yearly averages cannot describe peak congestion. The 4.5% growth rate and intervention mechanics are scenario assumptions. First High year is a conditional threshold crossing, not a certified road-failure date or a policy instruction.")
    doc.save(OUT / "working_simulation_artifact.docx")


def code_pdf():
    target = OUT / "python_source_code.pdf"
    page_w, page_h = landscape(letter)
    c = canvas.Canvas(str(target), pagesize=(page_w, page_h))
    margin = 36
    font_size = 7.2
    line_h = 9.4
    page_num = 0
    y = 0
    name = ""

    def page(file_name):
        nonlocal y, page_num, name
        if page_num:
            c.showPage()
        page_num += 1
        name = file_name
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin, page_h - margin, "Metro Manila AADT simulation - Python source code")
        c.setFont("Helvetica", 8)
        c.drawRightString(page_w - margin, page_h - margin, f"{name} | page {page_num}")
        c.line(margin, page_h - margin - 7, page_w - margin, page_h - margin - 7)
        y = page_h - margin - 24

    for source in (ROOT / "aadt_model.py", ROOT / "run_simulation.py"):
        page(source.name)
        for number, raw in enumerate(source.read_text(encoding="utf-8").splitlines(), 1):
            # Preserve indentation. Long source lines continue with a blank line number.
            prefix = f"{number:>4}  "
            available = page_w - 2 * margin - stringWidth(prefix, "Courier", font_size)
            max_chars = int(available / stringWidth("M", "Courier", font_size))
            chunks = [raw[i:i + max_chars] for i in range(0, len(raw), max_chars)] or [""]
            for k, chunk in enumerate(chunks):
                if y < margin + 10:
                    page(source.name)
                c.setFont("Courier", font_size)
                c.drawString(margin, y, (prefix if k == 0 else "      ") + chunk.encode("ascii", "replace").decode("ascii"))
                y -= line_h
    c.save()


def copy_submission_inputs():
    from shutil import copy2
    pairs = [
        (ROOT / "aadt_model.py", OUT / "aadt_model.py"),
        (ROOT / "run_simulation.py", OUT / "run_simulation.py"),
        (ROOT / "requirements-cli.txt", OUT / "requirements-cli.txt"),
        (ROOT / "README.md", OUT / "README.md"),
        (ROOT / "data" / "metro_manila_aadt_simulation_input_2020_2025.csv", OUT / "metro_manila_aadt_simulation_input_2020_2025.csv"),
        (ROOT / "outputs" / "road_year_forecast.csv", OUT / "sample_road_year_forecast.csv"),
        (ROOT / "outputs" / "scenario_2035_summary.csv", OUT / "sample_2035_summary.csv"),
    ]
    for src, dst in pairs:
        copy2(src, dst)


if __name__ == "__main__":
    project_paper()
    artifact_guide()
    code_pdf()
    copy_submission_inputs()
    print("Created submission files in", OUT)
